# -*- coding: utf-8 -*-
"""词条 xlsx 导出 / 译文 xlsx 导回。

导出用 xlsxwriter（C 级写入）：实测 17050 行带样式 1.4s，比 openpyxl
（逐格样式设置，约 200s）快两个数量级；产物与 openpyxl 双向兼容。
导回用 openpyxl read_only（流式，1.2s，已足够快）。

导出列约定
----------
主表「词条」（翻译对象，导回只读这张表）：
A  词条ID   —— 词条哈希（SHA1 前 12 位），导回时仅作审计/防错配，不参与匹配
B  类型     —— Description | DropDown | Description+DropDown（模板行为族内并集）
C  源文本   —— 英文原文（唯一键）；数字模板行形如 'Huw +{0}'（{0} = 数字占位）
D  译文     —— 空待填；术语表命中时自动预填（可改）
E  出现次数 —— 该文本在 CT 中的总出现次数（模板行为族内合计）
F  分类     —— normal（需翻译）| param（疑似参数/字段名，有术语表译文）
G  备注     —— 术语表自动译文 / 需人工翻译 / 模板合并说明
H  上下文   —— 父级描述链（导入不读取，消歧参考）

副表「参数保留」（slot25/sfxld_170 这类纯标识符，无术语表译文、
翻译纯属浪费）——不进主表，不参与导回，仅留审计；CT 中原样保留。

数字模板族
----------
'Black Knight Captain Huw +1'..'+10' 这类仅数字不同的变体合并为一行
'Huw +{0}'，翻译一次；apply 对 CT 中每个具体变体按出现顺序回填数字。
译文必须保留与源文本等量的 {0} 占位符（否则该变体跳过并报告）。

导回规则
--------
- 仅读主表「词条」；译文列非空的行生效；空译文 = 不替换。
- A 列词条 ID 必须与源文本哈希一致（防止用户改源列造成错配）；
  不一致的行忽略并报告。删除任意行均安全（apply 按源文本键匹配，
  与行序/行数零耦合）。
"""
from __future__ import annotations

import hashlib
import re
from pathlib import Path

import xlsxwriter
from openpyxl import load_workbook

HEADERS = ["词条ID", "类型", "源文本", "译文", "出现次数", "分类", "备注", "上下文"]
COL_ID, COL_TYPE, COL_SRC, COL_DST, COL_COUNT, COL_CLASS, COL_NOTE = range(7)
COL_CTX = 7   # 上下文列（导入时不读取，仅校对/消歧参考）

_WIDTHS = [14, 18, 60, 40, 10, 10, 34, 26]

_HAS_NUM = re.compile(r"\d|%|\+|#|\[|\]")
# CJK 混排源检测：含假名 = 日文混排（CT 作者注记，建议汉化）；
# 仅含汉字无假名 = 已是中文（再"翻译"有改坏风险，建议保留）
_KANA = re.compile(r"[\u3040-\u30ff]")
_HAN = re.compile(r"[\u4e00-\u9fff]")


def term_id(text: str) -> str:
    return hashlib.sha1(text.encode("utf-8")).hexdigest()[:12]


# 数字模板：数字跑 → {0} 占位（'Huw +1'/'Huw +10' → 'Huw +{0}'）
_DIG_RUN = re.compile(r"\d+")
_TPL_PH = "{0}"


def make_template(text: str) -> str:
    """把文本中所有数字跑替换为 {0} 占位符。"""
    return _DIG_RUN.sub(_TPL_PH, text)


def fill_template(template_dst: str, digit_runs: list[str]) -> str | None:
    """按顺序把译文中的 {0} 占位符回填为数字跑；数量不符返回 None（宁可漏替换，不错替换）。"""
    ph = template_dst.count(_TPL_PH)
    if ph == 0 or ph != len(digit_runs):
        return None
    parts = template_dst.split(_TPL_PH)
    return parts[0] + "".join(r + p for r, p in zip(digit_runs, parts[1:]))


def merge_numeric_families(terms: list[dict],
                           auto: dict[str, str]) -> tuple[list[dict], dict]:
    """数字模板族合并：'Black Knight Captain Huw +1'..'+10' 只导出一行
    'Black Knight Captain Huw +{0}'，翻译一次，apply 按数字回填展开。

    合并条件（保守）：同一模板且同一分类、族内 ≥2 条变体、
    族内没有任何一条命中术语表预填（有预填的变体保持原样，不损失精度）。
    """
    families: dict[tuple[str, str], list[dict]] = {}
    for t in terms:
        src = t["source"]
        if _TPL_PH in src or not _DIG_RUN.search(src):
            continue
        families.setdefault((make_template(src), t["class"]), []).append(t)

    merged_keys: set[int] = set()
    merged_rows: list[dict] = []
    for (tpl, cls), members in families.items():
        if len(members) < 2 or any(m["source"] in auto for m in members):
            continue
        head = max(members, key=lambda m: m["count"])
        row = dict(head)
        row["source"] = tpl
        row["count"] = sum(m["count"] for m in members)
        row["type"] = sorted({x for m in members for x in m["type"]})
        row["entries"] = sorted({x for m in members for x in m["entries"]})
        ctxs: list[str] = []
        for m in members:
            for c in m.get("contexts", []):
                if c not in ctxs:
                    ctxs.append(c)
        row["contexts"] = ctxs[:4]
        row["variants"] = len(members)
        merged_rows.append(row)
        merged_keys.update(id(m) for m in members)

    out = [t for t in terms if id(t) not in merged_keys] + merged_rows
    out.sort(key=lambda r: (-r["count"], r["source"]))
    info = {"families": len(merged_rows),
            "variants_merged": sum(r["variants"] - 1 for r in merged_rows)}
    return out, info


def export_terms_xlsx(terms: list[dict], out_path: str | Path,
                      auto_translations: dict[str, str] | None = None,
                      partial_notes: dict[str, str] | None = None,
                      include_param: bool = True,
                      contexts: dict[str, list[str]] | None = None,
                      allowed_types: set[str] | None = None,
                      merge_families: bool = True) -> dict:
    """导出词条表（xlsxwriter）。

    terms: extract.extract_terms() 的 result["terms"]（已按出现次数降序）
    auto_translations: 术语表整串精确命中 {source: zh}（直接填译文列）
    partial_notes: 短语级词表建议 {source: 建议译文}（只进备注列，
        供翻译软件参考，不自动填入译文列，避免中英混合文本）
    include_param: False 时跳过 param 类（只导出需翻译的 normal）
    contexts: {source: [父级描述...]} 场景消歧参考（H 列，导入不读取）
    allowed_types: 类型开关（None = 全部）；行类型与开关无交集则不导出
    merge_families: True 时内部做数字模板族合并；调用方已自行合并
        （需要先于 partial 建议计算）则传 False
    """
    wb = xlsxwriter.Workbook(str(Path(out_path).resolve()))
    ws = wb.add_worksheet("词条")
    ws_aux = wb.add_worksheet("参数保留")
    fmt_head = wb.add_format({"bold": True, "bg_color": "#D9E1F2", "align": "center"})
    fmt_auto = wb.add_format({"bg_color": "#E2EFDA"})    # 术语表预填（浅绿）

    auto = auto_translations or {}
    partial = partial_notes or {}
    ctx_map = contexts or {}
    # 数字模板族合并：仅数字不同的变体（+1/+2/…）导出一行模板，
    # 翻译一次，apply 按数字回填。只合并无预填的变体（有预填保持原样）。
    if merge_families:
        terms, fam_info = merge_numeric_families(terms, auto)
    else:
        fam_info = {}
    n_normal = n_auto = n_param_kept = 0
    rows = []
    aux_rows = []
    for t in terms:
        if allowed_types is not None and not (set(t["type"]) & allowed_types):
            continue
        if t["class"] == "param" and not include_param:
            continue
        src = t["source"]
        dst = auto.get(src, "")
        if dst:
            note, fmt = "术语表自动译文（可修改）", fmt_auto
            n_auto += 1
        elif t["class"] == "param":
            # 纯标识符/符号/代码成分且无术语表译文：翻译纯属浪费，
            # 挪到副表留审计（CT 中原样保留），主表只留真正需要翻译的内容
            aux_rows.append((term_id(src), "+".join(t["type"]), src))
            n_param_kept += 1
            continue
        else:
            note, fmt = "需人工翻译", None
            if _TPL_PH in src:
                note = (f"数字模板：{t.get('variants', '?')} 条仅数字不同的变体已合并，"
                        "翻译一次；译文保留 {0} 占位符，数字回填自动完成")
            elif _HAS_NUM.search(src):
                note += "；含数字/符号，翻译时保留"
            if _KANA.search(src):
                note += "；源含日文，建议汉化"
            elif _HAN.search(src):
                note += "；源已含中文，建议保留原样或仅校对"
            if src in partial:
                note += f"；词表建议参考：{partial[src]}"
            n_normal += 1
        ctx = " > ".join(ctx_map.get(src, [])[:2])
        rows.append((term_id(src), "+".join(t["type"]), src, dst,
                     t["count"], t["class"], note, fmt, ctx))

    ws.write_row(0, 0, HEADERS, fmt_head)
    for i, r in enumerate(rows, start=1):
        # 源文本/词条ID 强制字符串：数字型源（"128"、"1.5"）被 xlsx 识别为
        # 数值会破坏导回哈希校验（实测 issues:1），write_string 消除该边界
        ws.write_string(i, COL_ID, r[0])
        ws.write_string(i, COL_TYPE, r[1])
        ws.write_string(i, COL_SRC, r[2])
        if r[3]:
            ws.write_string(i, COL_DST, r[3])
        else:
            ws.write_blank(i, COL_DST, None)
        ws.write_number(i, COL_COUNT, r[4])
        ws.write_string(i, COL_CLASS, r[5])
        ws.write_string(i, COL_NOTE, r[6])
        if r[8]:
            ws.write_string(i, COL_CTX, r[8])
    for c, w in enumerate(_WIDTHS):
        ws.set_column(c, c, w)
    ws.freeze_panes(1, 0)
    ws.autofilter(0, 0, len(rows), len(HEADERS) - 1)

    # 副表：参数保留（仅审计，不导回）
    ws_aux.write_row(0, 0, ["词条ID", "类型", "源文本"], fmt_head)
    for i, r in enumerate(aux_rows, start=1):
        ws_aux.write_string(i, 0, r[0])
        ws_aux.write_string(i, 1, r[1])
        ws_aux.write_string(i, 2, r[2])
    for c, w in enumerate([14, 18, 60]):
        ws_aux.set_column(c, c, w)
    ws_aux.freeze_panes(1, 0)

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.close()
    return {"rows": len(rows), "auto": n_auto, "param_kept": n_param_kept,
            "manual": n_normal, **fam_info}


def import_translations_xlsx(xlsx_path: str | Path) -> tuple[dict[str, str], list[str], set[str]]:
    """导回译文。

    返回 (translations, issues, exported)：
    translations: {源文本: 译文}，空译文跳过
    issues:       问题列表（源列被改动导致哈希不一致等，相关行已忽略）
    exported:     导出文件中出现的全部源文本集合（含空译文行，用于界定 fallback）
    """
    wb = load_workbook(str(xlsx_path), read_only=True, data_only=True)
    ws = wb["词条"] if "词条" in wb.sheetnames else wb.active

    translations: dict[str, str] = {}
    exported: set[str] = set()
    issues: list[str] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None or len(row) < 4:
            continue
        src = row[COL_SRC]
        if src is None or str(src).strip() == "":
            continue
        src = str(src)
        exported.add(src)
        if row[COL_ID] and str(row[COL_ID]).strip() != term_id(src):
            issues.append(f"词条ID与源文本不匹配，已忽略该行: {src[:60]!r}")
            continue
        dst = row[COL_DST]
        if dst is None or str(dst).strip() == "":
            continue
        dst = str(dst).strip()
        if src in translations and translations[src] != dst:
            issues.append(f"源文本重复且译文冲突: {src[:60]!r}")
        translations[src] = dst
    wb.close()
    return translations, issues, exported
