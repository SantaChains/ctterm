# -*- coding: utf-8 -*-
"""增量翻译：翻译记忆（TM）构建、delta 导出、译文集合并。

核心原则
--------
apply 永远从原始 EN CT 全量重建，绝不打补丁到已译 CT（已译文本按
源文本键查不到，行级匹配必然 missed）；增量只发生在"译文集"层面。

三层继承
--------
1. TM 层：新表源文本命中翻译记忆 → 直接继承旧译文，零成本。
   TM 来源：旧 translations.json + 旧双语表对齐（--en/--zh 同 ID 对齐，
   en_text != zh_text 的键值对即 TM，1585 个共享 Description 全部受益）。
2. ID 层：--en/--zh 对齐时，同 (entry_id, field) 键但源文本已变的条目
   → 旧中文写进 delta.xlsx 备注列作"ID 继承参考"，人工/AI 校订。
3. 词表层：术语表整串命中 + partial 建议（与 export 相同）。

manifest 状态文件（借鉴 obsidian-i18n）
---------------------------------------
每次 delta 落盘 <out>.manifest.json：输入路径、TM 命中统计、时间戳，
供后续审计与 GUI 状态恢复。
"""
from __future__ import annotations

import json
import time
from pathlib import Path

from .ctmodel import CTModel
from .extract import extract_terms
from .glossary import is_usable_zh
from .xlsxio import export_terms_xlsx


def build_tm(tm_json: str | Path | None,
             en_ct: str | Path | None,
             zh_ct: str | Path | None) -> tuple[dict[str, str], dict[str, tuple[str, str]]]:
    """构建翻译记忆。返回 (tm, carry)。

    tm:    {en_text: zh_text} 源文本级翻译记忆（可直接继承）。
           仅收录 en→zh 文本不同（zh 确为译文）的键值对。
    carry: {entry_id: (old_en_text, old_zh_text)} 按 ID 对齐的继承参考。
           新版源文本未命中 TM 时，若其词条 ID 在旧表存在且旧 EN 文本
           与新源文本不同（跨版本改词），旧译文作为"ID继承参考"写备注，
           供人工/AI 校订——绝不直接进替换字典（旧译对应的是旧文本）。
    """
    tm: dict[str, str] = {}
    if tm_json:
        p = Path(tm_json)
        if p.exists():
            data = json.loads(p.read_text(encoding="utf-8"))
            tm.update((k, v) for k, v in data.items()
                      if isinstance(v, str) and v and is_usable_zh(v))
    carry: dict[str, tuple[str, str]] = {}
    if en_ct and zh_ct:
        snap_en = _text_snapshot(en_ct)
        snap_zh = _text_snapshot(zh_ct)
        for key, en_text in snap_en.items():
            zh_text = snap_zh.get(key)
            if zh_text and zh_text != en_text and is_usable_zh(zh_text):
                # 同 ID 键下 en != zh：zh 即 en 的既有译文
                if en_text not in tm:
                    tm[en_text] = zh_text
                carry.setdefault(key[0], (en_text, zh_text))
    return tm, carry


def _text_snapshot(path: str | Path) -> dict[tuple[str, str], str]:
    m = CTModel(path)
    snap: dict[tuple[str, str], str] = {}
    for e in m.entries:
        if e.description is not None:
            snap[(e.entry_id, "Description")] = e.description
        for it in e.dropdown:
            if it.name:
                snap[(e.entry_id, f"DropDown:{it.item_id}")] = it.name
    return snap


def run_delta(new_ct: str, out_xlsx: str,
              zh_ct: str | None = None,
              en_ct: str | None = None,
              tm_json: str | None = None,
              glossary: bool = False,
              user_terms: str | None = None,
              conflicts: bool = False,
              prefill_partial: bool = False) -> dict:
    """增量导出：只处理新表词条，TM 命中预填，含上下文列与 manifest。"""
    from .cli import _load_glossary  # 局部导入避免环

    result = extract_terms(new_ct)          # 不落盘 terms.json（词条随 xlsx 走）
    terms = result["terms"]
    stats = dict(result["stats"])

    tm, carry = build_tm(tm_json, en_ct, zh_ct)

    auto: dict[str, str] = {}
    notes_extra: dict[str, str] = {}
    for t in terms:
        src = t["source"]
        dst = tm.get(src)
        if dst and dst != src and is_usable_zh(dst):
            auto[src] = dst
        else:
            # 源文本未直接命中 TM：按词条 ID 回溯旧表——
            # 同 ID 的旧 EN 文本与当前源不同（跨版本改词）时，
            # 旧译文仅作备注参考（对应旧文本，不可信为现文本译文）
            for eid in t["entries"]:
                pair = carry.get(eid)
                if pair and pair[0] != src and pair[1] != src:
                    notes_extra.setdefault(
                        src, f"ID继承参考：{pair[1]}")
                    break

    # 模板族合并先于 partial：模板行也能拿到词表短语建议
    from .xlsxio import merge_numeric_families
    terms, fam_info = merge_numeric_families(terms, auto)
    contexts = {t["source"]: t["contexts"] for t in terms if t.get("contexts")}

    partial: dict[str, str] = {}
    if glossary:
        g = _load_glossary(user_terms, conflicts)
        exact = dict(g.items())
        for t in terms:
            src = t["source"]
            if src in auto:
                continue
            dst = exact.get(src)
            if dst and dst != src and is_usable_zh(dst):
                auto[src] = dst
        if prefill_partial:
            from .matcher import LongestMatcher
            matcher = LongestMatcher(list(g.filtered_items()))
            for t in terms:
                src = t["source"]
                if src in auto or t["class"] != "normal":
                    continue
                suggested, hits = matcher.replace_with_hits(src)
                if hits and is_usable_zh(suggested):
                    partial[src] = suggested

    # 备注列合并：术语表建议 + ID 继承参考
    merged_terms = [dict(t) for t in terms]
    info = export_terms_xlsx(merged_terms, out_xlsx,
                             auto_translations=auto,
                             partial_notes=partial,
                             include_param=True,
                             contexts=contexts,
                             merge_families=False)
    info.update(fam_info)
    if notes_extra:
        _append_notes(out_xlsx, notes_extra)

    manifest = {
        "kind": "eldenct-delta",
        "new_ct": str(new_ct),
        "zh_ct": zh_ct, "en_ct": en_ct, "tm_json": tm_json,
        "tm_size": len(tm), "carry_size": len(carry),
        "prefilled": len(auto),
        "stats": stats,
        "created": time.strftime("%Y-%m-%d %H:%M:%S"),
    }
    Path(out_xlsx).with_suffix(".xlsx.manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=1), encoding="utf-8")
    return {"xlsx": str(out_xlsx), "tm_prefilled": len(auto),
            "carry_refs": len(notes_extra), "tm_size": len(tm), **stats}


def _append_notes(xlsx_path: str, notes: dict[str, str]):
    """向已生成 xlsx 的备注列追加 ID 继承参考（openpyxl 原位改写）。"""
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path)
    ws = wb["词条"]
    hits = 0
    for row in ws.iter_rows(min_row=2):
        src = row[2].value
        if src in notes:
            old = row[6].value or ""
            row[6].value = (old + "；" if old else "") + notes[src]
            hits += 1
    wb.save(xlsx_path)
    return hits


def run_merge(base_json: str, delta_json: str, out_json: str) -> dict:
    """合并译文集：delta 优先；冲突（同键异值）以 delta 为准并报告。"""
    base = json.loads(Path(base_json).read_text(encoding="utf-8"))
    delta = json.loads(Path(delta_json).read_text(encoding="utf-8"))
    conflicts = [k for k, v in delta.items() if k in base and base[k] != v]
    merged = dict(base)
    merged.update(delta)
    Path(out_json).write_text(
        json.dumps(merged, ensure_ascii=False, indent=1), encoding="utf-8")
    return {"merged": len(merged),
            "base": len(base), "delta": len(delta),
            "conflicts_delta_wins": len(conflicts),
            "conflict_preview": conflicts[:10]}
