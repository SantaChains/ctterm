# -*- coding: utf-8 -*-
"""构建合格术语表 glossary_qualified.tsv/json。

原则（用户硬约束）：
  1. 键 100% 来自 CT 真实词条（CT 不会出现的键一律排除）
  2. 译文来源优先级：Hexinton ID List 官方多语言 xlsx > 杂乱/id列表.xlsx
     > 人工校对 ct英文校对版（entry 级整句）> combine_flat.json
  3. 防错过滤：译文须含 CJK 且无全角 ASCII；键间子串冲突审计报告
  4. 只生成新文件，绝不写任何源文件
"""
import json
import os
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

_REPO_ROOT = Path(__file__).resolve().parent.parent   # data/ → 仓库根
sys.path.insert(0, str(_REPO_ROOT))

from eldenct.ctmodel import CTModel

# 素材根目录：放 CT 源表、Hexinton ID List xlsx、社区词表等外部素材。
# 默认 <仓库根>/eldenring/（与 ELDENCT_WORD_ROOT 同一套布局），可传参或环境变量覆盖
ROOT = Path(sys.argv[1] if len(sys.argv) > 1
            else os.environ.get("ELDENCT_WORD_ROOT") or (_REPO_ROOT / "eldenring"))
CT = ROOT / "Hexinton-v8.0.4.CT"
OUT_TSV = ROOT / "glossary_qualified.tsv"
OUT_JSON = ROOT / "glossary_qualified.json"
OUT_MISS = ROOT / "glossary_qualified_missing.tsv"

# ---------- 1. CT 真实词条 ----------
model = CTModel(str(CT))
# 词条来源分类：Description 词条可用 entry 级人工校对（校对译文是 Description 的译文）；
# DropDown 名称绝不能用 entry 级译文（那是 Description 的译文，强行配对会错位，
# 如地名 'Weeping Peninsula - Tombsward Cave' 被错配成功能注释"恩惠下拉菜单"）。
term_desc = {}   # 文本 -> entry_ids（仅 Description）
term_dd = {}     # 文本 -> entry_ids（仅 DropDown 名称）
for e in model.entries:
    for it in e.dropdown:
        t = it.name.strip()
        if t:
            term_dd.setdefault(t, set()).add(e.entry_id)
    d = e.description
    if d:
        term_desc.setdefault(d, set()).add(e.entry_id)
# 与引擎键一致：description 走引号剥离（首尾都引号才剥）
def strip_desc_quotes(s: str) -> str:
    if len(s) >= 2 and s[0] == '"' and s[-1] == '"':
        return s[1:-1]
    return s
final_terms = {}
term_src_kind = {}   # 文本 -> {'desc': bool, 'dd': bool}
for t, ids in term_desc.items():
    t2 = strip_desc_quotes(t)
    final_terms.setdefault(t2, set()).update(ids)
    term_src_kind.setdefault(t2, {"desc": False, "dd": False})["desc"] = True
for t, ids in term_dd.items():
    t2 = strip_desc_quotes(t)
    final_terms.setdefault(t2, set()).update(ids)
    term_src_kind.setdefault(t2, {"desc": False, "dd": False})["dd"] = True
# 过滤：纯数字/纯符号/空，以及已含中文/日文的双语或中文键（非 EN→ZH 翻译对象）
def is_junk(s: str) -> bool:
    if not s.strip():
        return True
    if re.search(r"[\u4e00-\u9fff]", s):   # 键内已含 CJK（如 "英文 | 日文" 双语说明）
        return True
    return not re.search(r"[A-Za-z]", s)
junk = [t for t in final_terms if is_junk(t)]
for t in junk:
    del final_terms[t]
print(f"CT 真实词条（过滤垃圾后）: {len(final_terms)}  | 剔除纯数字/纯符号 {len(junk)}")

# ---------- 2. 译文源 ----------
# 2a. Hexinton ID List xlsx（官方多语言）：English -> Chinese
def load_xlsx_en_zh(path, en_cols, zh_cols):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    d = {}
    for sn in wb.sheetnames:
        ws = wb[sn]
        header = None
        for row in ws.iter_rows(values_only=True):
            if header is None:
                if not row or all(v is None for v in row):
                    continue
                header = [str(v).strip().lower() if v else "" for v in row]
                continue
            if not row or all(v is None for v in row):
                continue
            en = zh = None
            for i, h in enumerate(header):
                if i >= len(row):
                    break
                v = row[i]
                if v is None:
                    continue
                if any(c in h for c in en_cols):
                    if en is None:
                        en = str(v).strip()
                if any(c in h for c in zh_cols):
                    if zh is None:
                        zh = str(v).strip()
            if en and zh:
                d[en] = zh
    wb.close()
    return d

print("读 Hexinton ID List xlsx ...")
src_official = load_xlsx_en_zh(ROOT / "Hexinton ID List 3.0 DLC.xlsx",
                               ["english", "name"], ["chinese", "描述 (中文"])
print("  官方 xlsx EN->ZH:", len(src_official))

print("读 杂乱/id列表.xlsx ...")
src_idlist = load_xlsx_en_zh(ROOT / "杂乱" / "id列表.xlsx",
                             ["english", "name"], ["chinese", "描述 (中文"])
print("  id列表 xlsx EN->ZH:", len(src_idlist))

# 2b. combine_flat.json（兜底词表）
src_flat = json.loads((ROOT / "combine_flat.json").read_text(encoding="utf-8"))
print("  combine_flat EN->ZH:", len(src_flat))

# 2c. 人工校对版：entry_id(十进制) -> 中文，反查 entry 英文
print("读 杂乱/ct英文校对版.txt ...")
src_manual = {}
for line in (ROOT / "杂乱" / "ct英文校对版.txt").read_text(encoding="utf-8").splitlines():
    line = line.strip()
    if not line or "," not in line:
        continue
    k, v = line.split(",", 1)
    v = v.strip()
    if v:
        src_manual[k.strip()] = v
print("  校对版 ID->ZH:", len(src_manual))
# CT entry_id（十进制字符串，与校对 key 同体系）
manual_hit = sum(1 for e in model.entries if e.entry_id in src_manual)
print(f"  校对版与 CT entry 匹配: {manual_hit}/{len(model.entries)}")

# ---------- 3. 逐词条取译 ----------
# 中文数字（0-99，用于 Lv2→二级 这类合法转换的校验）
_CN = ["零", "一", "二", "三", "四", "五", "六", "七", "八", "九"]
def _cn_num(n: int) -> str:
    if n <= 9:
        return _CN[n]
    if n < 20:
        return "十" + (_CN[n - 10] if n > 10 else "")
    return _CN[n // 10] + "十" + (_CN[n % 10] if n % 10 else "")

def num_ok(k: str, v: str) -> bool:
    """数字保真：原文的小数数字必须原样出现在译文；
    整数数字须以阿拉伯数字或中文数字形式出现（Lv2→二级 通过，v2.2→v2.1 拒绝）。"""
    for m in re.finditer(r"\d+(?:\.\d+)?", k):
        tok = m.group(0)
        if "." in tok:
            if tok not in v:
                return False
        else:
            n = int(tok)
            if tok in v or (0 <= n <= 99 and _cn_num(n) in v):
                continue
            return False
    return True

def zh_ok(v: str) -> bool:
    if not v:
        return False
    if not re.search(r"[\u4e00-\u9fff]", v):
        return False
    # 全角 ASCII 视为误译信号
    if re.search(r"[\uff01-\uff5e]", v):
        return False
    return True

def lookup(t: str):
    """按优先级取译文；返回 (译文, 来源) 或 (None, None)"""
    if t in src_official and zh_ok(src_official[t]) and num_ok(t, src_official[t]):
        return src_official[t], "official_xlsx"
    if t in src_idlist and zh_ok(src_idlist[t]) and num_ok(t, src_idlist[t]):
        return src_idlist[t], "idlist_xlsx"
    # 校对版：仅限该文本确为某 entry 的 Description 时可用（entry 级译文 ≠ DD 名称译文）
    if term_src_kind.get(t, {}).get("desc"):
        manual_z = set()
        for eid in final_terms[t]:
            if eid in src_manual:
                manual_z.add(src_manual[eid])
        if len(manual_z) == 1:
            v = next(iter(manual_z))
            if zh_ok(v) and num_ok(t, v):
                return v, "manual"
    if t in src_flat and zh_ok(src_flat[t]) and num_ok(t, src_flat[t]):
        return src_flat[t], "combine_flat"
    # 不区分大小写兜底（限 combine_flat）
    tl = t.lower()
    for k, v in src_flat.items():
        if k.lower() == tl and zh_ok(v) and num_ok(t, v):
            return v, "combine_flat_ci"
    return None, None

result = {}
src_stat = Counter()
missing = []
for i, t in enumerate(sorted(final_terms)):
    v, s = lookup(t)
    if v:
        result[t] = v
        src_stat[s] += 1
    else:
        missing.append(t)

print()
print("=== 结果 ===")
print("有译文:", len(result), "| 缺译文:", len(missing))
print("来源分布:", dict(src_stat))

# ---------- 4. 键间冲突审计（子串 + 整词边界） ----------
keys = sorted(result, key=len, reverse=True)
short_keys = [k for k in keys if len(k) < len(max(keys, key=len))]
sub_conflicts = []
for long_k in keys:
    for short_k in keys:
        if short_k == long_k or len(short_k) >= len(long_k):
            continue
        if short_k in long_k:
            # 检查是否整词边界（前后非字母数字）——整词子串无碍（长键先匹配）
            idx = 0
            while True:
                idx = long_k.find(short_k, idx)
                if idx < 0:
                    break
                before = long_k[idx - 1] if idx > 0 else ""
                after = long_k[idx + len(short_k)] if idx + len(short_k) < len(long_k) else ""
                if before and before.isalnum() and after and after.isalnum():
                    sub_conflicts.append((short_k, long_k))
                idx += 1
print()
print("子串冲突审计（短键嵌于长键内部、两侧均紧贴字母数字）:", len(sub_conflicts))
for s, l in sub_conflicts[:12]:
    print(f"  {s!r} 内嵌于 {l!r}")

# ---------- 5. 写出 ----------
with OUT_TSV.open("w", encoding="utf-8", newline="") as f:
    for t in sorted(result):
        f.write(f"{t}\t{result[t]}\n")
with OUT_JSON.open("w", encoding="utf-8") as f:
    json.dump(result, f, ensure_ascii=False, indent=1)
with OUT_MISS.open("w", encoding="utf-8") as f:
    for t in sorted(missing):
        f.write(t + "\n")
print()
print("写出:", OUT_TSV.name, len(result), "行")
print("写出:", OUT_JSON.name)
print("缺译文清单:", OUT_MISS.name, len(missing), "行")
